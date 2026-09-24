import streamlit as st
import requests
import io
from docx import Document
from docx.shared import Pt
from fpdf import FPDF

# Setup page configuration
st.set_page_config(page_title="LegalEase", layout="centered")

# Header
st.markdown("<h2 style='text-align: center;'>⚖️ LegalEase</h2>", unsafe_allow_html=True)
st.markdown("<h4 style='text-align: center;'>AI Legal Document Generator</h4>", unsafe_allow_html=True)
st.markdown("---")

# User Inputs
document_type = st.text_input("Document Type (Ex: Agreement, Contract, NDA)", "Freelance Work Contract")
parties = st.text_area("Parties Involved", "Jane Doe (Service Provider), TechNova Inc. (Client)")
terms = st.text_area("Terms & Conditions (Use semicolons for bullet points)", "Work must be delivered by May 15, 2025; Payment will be made within 7 days of invoice;")
dates = st.text_input("Effective Date", "April 15, 2025")

# Utility Functions
def sanitize_text(text):
    return text.replace("â€™", "'").replace("â€œ", '"').replace("â€", '"').replace("\r", "")

def format_docx(text, doc_type):
    doc = Document()
    doc.add_heading(doc_type, 0)
    for line in text.split('\n'):
        if line.strip():
            # Bold lines starting with #
            if line.startswith('#'):
                p = doc.add_paragraph()
                p.add_run(line.replace('#', '').strip()).bold = True
            else:
                doc.add_paragraph(line)
    
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

def format_pdf(text, doc_type):
    class PDF(FPDF):
        def header(self):
            self.set_font('Arial', 'B', 15)
            self.cell(0, 10, 'LegalEase', 0, 1, 'C')
            self.ln(10)

        def footer(self):
            self.set_y(-15)
            self.set_font('Arial', 'I', 8)
            self.cell(0, 10, f'Page {self.page_no()}', 0, 0, 'C')

    pdf = PDF()
    pdf.add_page()
    pdf.set_font("Arial", size=11)
    
    # Handle utf-8 encoding safely for FPDF
    safe_text = text.encode('latin-1', 'replace').decode('latin-1')
    
    # Title
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(0, 10, txt=doc_type, ln=1, align="C")
    pdf.ln(5)
    
    pdf.set_font("Arial", size=11)
    for line in safe_text.split('\n'):
        if line.startswith('#'):
            pdf.set_font("Arial", 'B', 11)
            pdf.multi_cell(0, 10, txt=line.replace('#', '').strip())
            pdf.set_font("Arial", size=11)
        else:
            pdf.multi_cell(0, 8, txt=line)
            
    # FPDF output as string (bytes)
    # Using 'S' dest returns a string (in latin-1), convert it to bytes for download
    try:
        pdf_bytes = pdf.output(dest='S').encode('latin-1')
    except Exception:
        pdf_bytes = pdf.output(dest='S')
        if type(pdf_bytes) == str:
            pdf_bytes = pdf_bytes.encode('latin-1')
    return pdf_bytes

def format_html_preview(text):
    html = "<div style='background-color: #1e1e1e; padding: 20px; border-radius: 10px; color: #fff; max-height: 500px; overflow-y: auto;'>"
    for line in text.split('\n'):
        if line.startswith('#'):
            html += f"<h4 style='color: #4da6ff;'>{line.replace('#', '').strip()}</h4>"
        elif line.strip():
            html += f"<p>{line}</p>"
    html += "</div>"
    return html

# Generate Document logic
if st.button("Generate Document", type="primary"):
    if not document_type or not parties or not terms or not dates:
        st.error("Please fill in all fields.")
    else:
        with st.spinner("Generating document... (This may take a few seconds)"):
            try:
                response = requests.post("http://localhost:8000/generate", json={
                    "document_type": document_type,
                    "parties": parties,
                    "terms": terms,
                    "dates": dates
                })
                
                if response.status_code == 200:
                    data = response.json()
                    if "error" in data:
                        st.error(f"Backend Error: {data['error']}")
                    else:
                        st.success("✅ Document Generated Successfully!")
                        generated_text = sanitize_text(data["document"])
                        
                        st.session_state['generated_text'] = generated_text
                        st.session_state['document_type'] = document_type
                else:
                    st.error(f"Failed to connect to backend: HTTP {response.status_code}")
            except requests.exceptions.ConnectionError:
                st.error("Error: Could not connect to the backend. Is the FastAPI server running on port 8000?")
            except Exception as e:
                st.error(f"An unexpected error occurred: {e}")

if 'generated_text' in st.session_state:
    st.markdown("### Preview")
    st.markdown(format_html_preview(st.session_state['generated_text']), unsafe_allow_html=True)
    
    st.markdown("---")
    
    if st.checkbox("✏️ Click to Edit Document"):
        st.session_state['generated_text'] = st.text_area("Edit Document Below:", st.session_state['generated_text'], height=300)
    
    # Downloads
    st.markdown("### Download")
    col1, col2, col3 = st.columns(3)
    file_prefix = st.session_state['document_type'].replace(' ', '_').lower()
    
    with col1:
        st.download_button("📄 Download as .TXT", data=st.session_state['generated_text'], file_name=f"{file_prefix}.txt")
    with col2:
        st.download_button("📝 Download as .DOCX", data=format_docx(st.session_state['generated_text'], st.session_state['document_type']), file_name=f"{file_prefix}.docx")
    with col3:
        st.download_button("📕 Download as .PDF", data=format_pdf(st.session_state['generated_text'], st.session_state['document_type']), file_name=f"{file_prefix}.pdf")
